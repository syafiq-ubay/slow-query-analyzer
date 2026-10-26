"""Streaming Excel export helpers."""

import os
from datetime import datetime, timezone
from typing import Any, Dict, Iterable, List, Optional
from zoneinfo import ZoneInfo


MYSQL_FIELDS = [
    "@timestamp",
    "user",
    "schema",
    "query_time",
    "lock_time",
    "rows_examined",
    "rows_sent",
    "query",
]

NGINX_FIELDS = [
    "@timestamp",
    "_log_name",
    "remote",
    "http_x_forwarded_for",
    "as_name",
    "method",
    "path",
    "query",
    "referer",
    "size",
    "rt",
    "agent",
    "code",
]

MERGED_FIELDS = [
    "mysql_@timestamp",
    "mysql_query_time",
    "mysql_schema",
    "mysql_user",
    # "mysql_host",
    # "mysql_query_type",
    "mysql_rows_examined",
    "mysql_rows_sent",
    "mysql_query",
    "http_timestamp",
    "http_method",
    "http_status_code",
    "http_response_time",
    "http_path",
    "http_query",
    "http_client_ip",
    "http_x_forwarded_for",
    "http_referer",
    "http_user_agent",
]


def export_to_excel(
    records: Iterable[Dict[str, Any]],
    output_path: str,
    fields: Optional[List[str]] = None,
    sheet_title: str = "Data",
    max_rows_per_sheet: int = 1_000_000,
) -> int:
    """Write records incrementally to an XLSX workbook with bounded memory."""
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required. Install it with: pip install openpyxl") from exc

    if max_rows_per_sheet <= 0:
        raise ValueError("max_rows_per_sheet must be greater than zero.")

    fieldnames = fields or MYSQL_FIELDS
    parent_dir = os.path.dirname(os.path.abspath(output_path))
    os.makedirs(parent_dir, exist_ok=True)

    workbook = Workbook(write_only=True)
    sheet_index = 1
    worksheet = workbook.create_sheet(title=sheet_title)
    worksheet.append(fieldnames)
    rows_in_sheet = 0
    total_written = 0

    for record in records:
        if rows_in_sheet >= max_rows_per_sheet:
            sheet_index += 1
            worksheet = workbook.create_sheet(title=f"{sheet_title} {sheet_index}")
            worksheet.append(fieldnames)
            rows_in_sheet = 0

        worksheet.append([
            _excel_value(_format_field_value(field, record.get(field, "")))
            for field in fieldnames
        ])
        rows_in_sheet += 1
        total_written += 1

    workbook.save(output_path)
    return total_written


def read_excel_records(
    input_path: str,
    required_fields: Optional[List[str]] = None,
) -> Iterable[Dict[str, Any]]:
    """Read records from all sheets of an XLSX workbook in read-only mode."""
    if not os.path.isfile(input_path):
        raise FileNotFoundError(f"Excel input file not found: {input_path}")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError("openpyxl is required. Install it with: pip install openpyxl") from exc

    workbook = load_workbook(input_path, read_only=True, data_only=True)
    found_headers = False
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            headers = next(rows, None)
            if not headers or all(header is None for header in headers):
                continue
            found_headers = True
            fieldnames = [str(header) if header is not None else "" for header in headers]
            missing_fields = set(required_fields or ()) - set(fieldnames)
            if missing_fields:
                missing = ", ".join(sorted(missing_fields))
                raise ValueError(
                    f"Excel sheet '{worksheet.title}' in '{input_path}' "
                    f"is missing required columns: {missing}"
                )
            for row in rows:
                if not any(value is not None for value in row):
                    continue
                yield dict(zip(fieldnames, row))
        if not found_headers:
            raise ValueError(f"No header row found in Excel file: {input_path}")
    finally:
        workbook.close()


def _excel_value(value: Any) -> Any:
    if isinstance(value, str) and len(value) > 32_000:
        return value[:32_000] + " ...[TRUNCATED]"
    return value


_TIMESTAMP_FIELDS = {"@timestamp", "mysql_@timestamp", "http_timestamp"}
_MONTHS = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_JAKARTA = ZoneInfo("Asia/Jakarta")


def _format_field_value(field: str, value: Any) -> Any:
    if field not in _TIMESTAMP_FIELDS:
        return value
    formatted = _format_timestamp(value)
    return formatted if formatted is not None else value


def _format_timestamp(value: Any) -> Optional[str]:
    """Display timestamps as `Oct 26, 2026 @ 08:56:50.238` in Jakarta time."""
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, (int, float)):
        epoch = float(value)
        if abs(epoch) >= 1_000_000_000_000:
            epoch /= 1000
        parsed = datetime.fromtimestamp(epoch, tz=timezone.utc)
    elif value:
        text = str(value).strip()
        try:
            parsed = datetime.fromisoformat(text.replace("Z", "+00:00"))
        except ValueError:
            return None
    else:
        return None

    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=_JAKARTA)
    parsed = parsed.astimezone(_JAKARTA)
    month = _MONTHS[parsed.month - 1]
    milliseconds = parsed.microsecond // 1000
    return (
        f"{month} {parsed.day:02d}, {parsed.year:04d} @ "
        f"{parsed.hour:02d}:{parsed.minute:02d}:{parsed.second:02d}.{milliseconds:03d}"
    )
